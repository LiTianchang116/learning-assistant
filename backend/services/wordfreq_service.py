import re
import os

from docx import Document
import pdfplumber
import openpyxl

from config import TEMP_DIR

STOPWORDS = {
    "the", "a", "an", "is", "are", "was", "were", "be", "been", "being",
    "have", "has", "had", "do", "does", "did", "will", "would", "could",
    "should", "may", "might", "shall", "can", "need", "dare", "ought",
    "used", "to", "of", "in", "for", "on", "with", "at", "by", "from",
    "as", "into", "through", "during", "before", "after", "above", "below",
    "between", "out", "off", "over", "under", "again", "further", "then",
    "once", "here", "there", "when", "where", "why", "how", "all", "both",
    "each", "few", "more", "most", "other", "some", "such", "no", "nor",
    "not", "only", "own", "same", "so", "than", "too", "very", "just",
    "because", "but", "and", "or", "if", "while", "about", "up", "down",
    "that", "this", "these", "those", "it", "its", "he", "she", "they",
    "them", "his", "her", "their", "my", "your", "we", "our", "i", "me",
    "you", "what", "which", "who", "whom", "whose", "am", "also", "much",
    "many", "well", "get", "got", "make", "made", "take", "come", "go",
    "know", "see", "think", "say", "give", "use", "find", "tell", "ask",
    "work", "seem", "feel", "try", "leave", "call", "one", "two", "first",
    "last", "long", "great", "little", "old", "big", "high", "different",
    "small", "large", "next", "early", "young", "important", "public",
    "bad", "same", "able", "new", "now", "way", "people", "time", "year",
    "day", "thing", "man", "world", "life", "hand", "part", "place",
    "case", "week", "company", "system", "program", "question", "government",
    "number", "night", "point", "home", "water", "room", "mother", "area",
    "money", "story", "fact", "month", "lot", "right", "study", "book",
    "eye", "job", "word", "business", "issue", "side", "kind", "head",
    "house", "service", "friend", "father", "power", "hour", "game", "line",
    "end", "member", "law", "car", "city", "community", "name", "president",
    "team", "minute", "idea", "body", "information", "back", "parent",
    "face", "others", "level", "office", "door", "health", "person",
    "art", "war", "history", "party", "result", "change", "morning",
    "reason", "research", "girl", "guy", "moment", "air", "teacher",
    "force", "education",
}


def extract_text_from_docx(path: str) -> str:
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def extract_text_from_pdf(path: str) -> str:
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                texts.append(text)
    return "\n".join(texts)


def extract_text_from_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return extract_text_from_docx(path)
    elif ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext == ".txt":
        return extract_text_from_txt(path)
    return ""


def count_words(text: str, use_stopwords: bool = True, min_freq: int = 1) -> list[dict]:
    words = re.findall(r"[a-zA-Z]+", text.lower())
    freq = {}
    for w in words:
        if len(w) < 2:
            continue
        if use_stopwords and w in STOPWORDS:
            continue
        freq[w] = freq.get(w, 0) + 1

    total = sum(freq.values())
    items = []
    for word, count in freq.items():
        if count < min_freq:
            continue
        items.append({
            "word": word,
            "count": count,
            "frequency": round(count / total * 100, 2) if total > 0 else 0,
        })
    items.sort(key=lambda x: x["count"], reverse=True)
    return items[:200]


def export_to_excel(items: list[dict]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "词频统计"
    ws.append(["排名", "单词", "出现次数", "频率(%)"])
    for i, item in enumerate(items, 1):
        ws.append([i, item["word"], item["count"], item["frequency"]])
    path = os.path.join(TEMP_DIR, "wordfreq_result.xlsx")
    wb.save(path)
    return path
